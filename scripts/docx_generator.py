"""
DOCX_GENERATOR
Renderiza dados JSON em arquivos DOCX formatados conforme padrão CJF.
Gera: Atas de Reunião e Minutas de Ofício.
"""

from datetime import datetime

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    print("[ERRO] python-docx não instalado. Execute: pip install python-docx")


MESES_PT_EXTENSO = {
    "01": "janeiro", "02": "fevereiro", "03": "março", "04": "abril",
    "05": "maio", "06": "junho", "07": "julho", "08": "agosto",
    "09": "setembro", "10": "outubro", "11": "novembro", "12": "dezembro"
}


def _data_por_extenso(data_str: str) -> str:
    """Converte 'DD/MM/YYYY' para 'DD de mês de YYYY'."""
    try:
        partes = data_str.split("/")
        if len(partes) == 3:
            dia, mes, ano = partes
            return f"{int(dia)} de {MESES_PT_EXTENSO.get(mes.zfill(2), mes)} de {ano}"
    except Exception:
        pass
    return data_str


class DocxGenerator:
    """Gera documentos DOCX no padrão CJF a partir de dados JSON."""

    def __init__(self, tipo_documento: str = "ata"):
        if not DOCX_AVAILABLE:
            raise RuntimeError("python-docx não está instalado.")
        self.doc = Document()
        self.tipo = tipo_documento
        self._configurar_pagina()
        self._configurar_estilos()

    def _configurar_pagina(self):
        from docx.shared import Cm
        section = self.doc.sections[0]
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)

    def _configurar_estilos(self):
        style = self.doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)
        pf = style.paragraph_format
        pf.line_spacing = Pt(16.5)  # ~1,5 linhas
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)

    # ------------------------------------------------------------------
    # CABEÇALHO INSTITUCIONAL
    # ------------------------------------------------------------------

    def _adicionar_cabecalho_institucional(self):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("CONSELHO DA JUSTIÇA FEDERAL")
        run.bold = True
        run.font.size = Pt(12)

        p2 = self.doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run("CENTRO NACIONAL DE INTELIGÊNCIA DA JUSTIÇA FEDERAL")
        run2.bold = True
        run2.font.size = Pt(11)

        self.doc.add_paragraph()

    # ------------------------------------------------------------------
    # ATA DE REUNIÃO
    # ------------------------------------------------------------------

    def adicionar_cabecalho_ata(self, metadata: dict):
        self._adicionar_cabecalho_institucional()

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("ATA DE REUNIÃO")
        run.bold = True
        run.font.size = Pt(13)

        self.doc.add_paragraph()

        # Bloco de informações
        info = self.doc.add_paragraph()
        info.add_run(f"Órgão: ").bold = True
        info.add_run(metadata.get("orgao", "Centro Nacional de Inteligência da Justiça Federal"))
        info.add_run("\n")
        info.add_run("Data: ").bold = True
        info.add_run(metadata.get("data", ""))
        info.add_run("   ")
        info.add_run("Horário: ").bold = True
        info.add_run(
            f"{metadata.get('horario_inicio', '')} às {metadata.get('horario_fim', '')}"
        )
        info.add_run("\n")
        info.add_run("Local: ").bold = True
        info.add_run(metadata.get("local", ""))
        info.add_run("\n")
        info.add_run("Tipo: ").bold = True
        info.add_run(metadata.get("tipo_reuniao", "Ordinária"))
        info.add_run("\n")
        info.add_run("Processo SEI: ").bold = True
        info.add_run(metadata.get("numero_sei", "[A PREENCHER]"))

        self.doc.add_paragraph()

    def adicionar_secao_participantes(self, participantes: list):
        h = self.doc.add_paragraph()
        run = h.add_run("1. PARTICIPANTES")
        run.bold = True
        run.font.size = Pt(12)

        for p in participantes:
            nome = p.get("nome", "")
            cargo = p.get("cargo", "")
            tribunal = p.get("tribunal", "")
            linha = f"– {nome}"
            if cargo:
                linha += f" ({cargo}"
                if tribunal:
                    linha += f" – {tribunal}"
                linha += ")"
            elif tribunal:
                linha += f" – {tribunal}"
            self.doc.add_paragraph(linha)

        self.doc.add_paragraph()

    def adicionar_secao_deliberacoes(self, temas: list):
        h = self.doc.add_paragraph()
        run = h.add_run("2. DELIBERAÇÕES")
        run.bold = True
        run.font.size = Pt(12)
        self.doc.add_paragraph()

        for i, tema in enumerate(temas, start=1):
            # Subtítulo do tema
            st = self.doc.add_paragraph()
            run_st = st.add_run(f"2.{i}. {tema.get('titulo', '')}")
            run_st.bold = True

            if tema.get("relator"):
                p = self.doc.add_paragraph()
                p.add_run("Relator(a): ").bold = True
                p.add_run(tema["relator"])

            if tema.get("resumo_apresentacao"):
                p = self.doc.add_paragraph()
                p.add_run("Resumo da Apresentação: ").bold = True
                p.add_run(tema["resumo_apresentacao"])

            # Discussões
            discussoes = tema.get("discussoes", [])
            if discussoes:
                p = self.doc.add_paragraph()
                p.add_run("Discussões:").bold = True
                for d in discussoes:
                    nome_d = d.get("participante", "")
                    cargo_d = d.get("cargo", "")
                    trib_d = d.get("tribunal", "")
                    resumo_d = d.get("resumo", "")
                    cabecalho = nome_d
                    if cargo_d:
                        cabecalho += f" ({cargo_d}"
                        if trib_d:
                            cabecalho += f" – {trib_d}"
                        cabecalho += ")"
                    par = self.doc.add_paragraph(style="List Bullet")
                    par.add_run(f"{cabecalho}: ").bold = bool(cabecalho)
                    par.add_run(resumo_d)

            if tema.get("deliberacao"):
                p = self.doc.add_paragraph()
                p.add_run("Deliberação: ").bold = True
                p.add_run(tema["deliberacao"])

            if tema.get("votacao"):
                p = self.doc.add_paragraph()
                p.add_run("Votação: ").bold = True
                p.add_run(tema["votacao"])

            self.doc.add_paragraph()

    def adicionar_secao_encaminhamentos(self, encaminhamentos: list):
        h = self.doc.add_paragraph()
        run = h.add_run("3. ENCAMINHAMENTOS")
        run.bold = True
        run.font.size = Pt(12)

        if not encaminhamentos:
            self.doc.add_paragraph("Nenhum encaminhamento registrado.")
            self.doc.add_paragraph()
            return

        table = self.doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"

        # Cabeçalho
        hdr = table.rows[0].cells
        hdr[0].text = "Descrição"
        hdr[1].text = "Responsável"
        hdr[2].text = "Prazo"
        hdr[3].text = "Status"
        for cell in hdr:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

        # Linhas
        for enc in encaminhamentos:
            row = table.add_row().cells
            row[0].text = enc.get("descricao", "")
            row[1].text = enc.get("responsavel", "")
            row[2].text = enc.get("prazo", "")
            row[3].text = enc.get("status", "Pendente")

        self.doc.add_paragraph()

    def adicionar_secao_normas(self, normas: list, julgados: list):
        if not normas and not julgados:
            return

        h = self.doc.add_paragraph()
        run = h.add_run("4. REFERÊNCIAS NORMATIVAS E JURISPRUDENCIAIS")
        run.bold = True
        run.font.size = Pt(12)

        if normas:
            self.doc.add_paragraph("Normas citadas:").runs[0].bold = True
            for n in normas:
                desc = n.get("descricao", "")
                texto = f"{n.get('tipo', '')} {n.get('numero', '')}"
                if desc:
                    texto += f" – {desc}"
                self.doc.add_paragraph(f"– {texto}")

        if julgados:
            p = self.doc.add_paragraph("Julgados citados:")
            p.runs[0].bold = True
            for j in julgados:
                tema_j = j.get("tema", "")
                texto = f"{j.get('tribunal', '')} – {j.get('numero', '')}"
                if tema_j:
                    texto += f" ({tema_j})"
                self.doc.add_paragraph(f"– {texto}")

        self.doc.add_paragraph()

    def adicionar_assinatura(self, coordenador: dict):
        self.doc.add_paragraph()
        self.doc.add_paragraph("_" * 50)
        self.doc.add_paragraph(coordenador.get("nome", "[Nome do Coordenador]"))

        cargo = coordenador.get("cargo", "")
        tribunal = coordenador.get("tribunal", "")
        if cargo:
            self.doc.add_paragraph(cargo)
        if tribunal:
            self.doc.add_paragraph(tribunal)

        data_hoje = datetime.now().strftime("%d/%m/%Y")
        p = self.doc.add_paragraph()
        p.add_run(f"Brasília, {_data_por_extenso(data_hoje)}")

    # ------------------------------------------------------------------
    # MINUTA DE OFÍCIO
    # ------------------------------------------------------------------

    def gerar_oficio(
        self,
        encaminhamento: dict,
        dados_ata: dict,
        tipo_oficio: str = "Encaminhamento",
        numero_oficio: str = "[A PREENCHER]"
    ):
        """Gera uma minuta de ofício a partir de um encaminhamento da ata."""
        ata = dados_ata.get("ata", {})
        metadata = ata.get("metadata", {})
        coordenador = metadata.get("coordenador", {})

        data_reuniao = metadata.get("data", "")
        data_extenso = _data_por_extenso(data_reuniao) if data_reuniao else "[DATA]"

        self._adicionar_cabecalho_institucional()

        # Data e número do ofício
        p_data = self.doc.add_paragraph()
        p_data.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_data.add_run(f"Brasília, {data_extenso}")

        self.doc.add_paragraph()

        p_num = self.doc.add_paragraph()
        p_num.add_run(f"Ofício nº {numero_oficio}/2026/CJF").bold = True

        self.doc.add_paragraph()

        # Destinatário
        responsavel = encaminhamento.get("responsavel", "[Destinatário]")
        p_dest = self.doc.add_paragraph()
        p_dest.add_run("A Sua Excelência o(a) Senhor(a)").bold = True
        p_dest.add_run(f"\n{responsavel}")
        p_dest.add_run("\n[Cargo do Destinatário]")
        p_dest.add_run("\n[Instituição/Tribunal]")
        p_dest.add_run("\n[Endereço]")

        self.doc.add_paragraph()

        # Assunto
        descricao = encaminhamento.get("descricao", "")
        p_assunto = self.doc.add_paragraph()
        p_assunto.add_run("Assunto: ").bold = True
        p_assunto.add_run(descricao[:120] if descricao else "[Assunto]")

        self.doc.add_paragraph()

        # Saudação
        self.doc.add_paragraph("Senhor(a) [Título],")
        self.doc.add_paragraph()

        # Corpo
        intro = (
            f"Cumprimentando-o(a) cordialmente, informo que o Centro Nacional de "
            f"Inteligência da Justiça Federal, instituído nos termos da Resolução CJF "
            f"nº 499, de 1º de outubro de 2018, com alterações pelas Resoluções CJF "
            f"nº 834/2023 e nº 944/2025, em reunião realizada em {data_extenso}, "
            f"deliberou o seguinte encaminhamento:"
        )
        self.doc.add_paragraph(intro)
        self.doc.add_paragraph()

        # Descrição do encaminhamento
        self.doc.add_paragraph(descricao)

        prazo = encaminhamento.get("prazo", "")
        if prazo and prazo != "A definir":
            self.doc.add_paragraph()
            p_prazo = self.doc.add_paragraph()
            p_prazo.add_run("Prazo para cumprimento: ").bold = True
            p_prazo.add_run(prazo)

        self.doc.add_paragraph()
        self.doc.add_paragraph(
            "Solicita-se, portanto, que sejam adotadas as providências necessárias "
            "ao cumprimento da deliberação acima, comunicando-se o resultado a este "
            "Centro Nacional de Inteligência."
        )

        self.doc.add_paragraph()
        self.doc.add_paragraph("Respeitosamente,")
        self.doc.add_paragraph()
        self.doc.add_paragraph()

        # Assinatura
        self.doc.add_paragraph("_" * 50)
        nome_coord = coordenador.get("nome", "Vânila Cardoso André de Moraes")
        self.doc.add_paragraph(nome_coord)
        self.doc.add_paragraph(
            coordenador.get("cargo", "Juíza Federal Auxiliar da Corregedoria-Geral")
        )
        self.doc.add_paragraph(
            "Coordenadora do Grupo Operacional do Centro Nacional de Inteligência "
            "da Justiça Federal"
        )

        data_hoje = datetime.now().strftime("%d/%m/%Y")
        hora_hoje = datetime.now().strftime("%H:%M")
        self.doc.add_paragraph(f"Data: {data_hoje}   Hora: {hora_hoje}")

    # ------------------------------------------------------------------
    # RELATÓRIO EXECUTIVO (estrutura básica)
    # ------------------------------------------------------------------

    def gerar_relatorio_executivo(self, lista_atas: list, periodo: str = ""):
        self._adicionar_cabecalho_institucional()

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("RELATÓRIO EXECUTIVO CONSOLIDADO")
        run.bold = True
        run.font.size = Pt(13)

        p2 = self.doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run(f"Atas do Grupo Operacional — Período: {periodo or '[Período]'}")

        self.doc.add_paragraph()

        total_reunioes = len(lista_atas)
        total_deliberacoes = sum(len(a.get("ata", {}).get("temas", [])) for a in lista_atas)
        total_encaminhamentos = sum(len(a.get("ata", {}).get("encaminhamentos", [])) for a in lista_atas)
        pendentes = sum(
            1 for a in lista_atas
            for e in a.get("ata", {}).get("encaminhamentos", [])
            if e.get("status", "Pendente") == "Pendente"
        )

        # 1. Executive Summary
        h1 = self.doc.add_paragraph()
        h1.add_run("1. EXECUTIVE SUMMARY").bold = True
        self.doc.add_paragraph(f"– Reuniões analisadas: {total_reunioes}")
        self.doc.add_paragraph(f"– Total de deliberações: {total_deliberacoes}")
        self.doc.add_paragraph(f"– Total de encaminhamentos: {total_encaminhamentos}")
        self.doc.add_paragraph(f"– Encaminhamentos pendentes: {pendentes}")
        self.doc.add_paragraph()

        # 2. Cronologia
        h2 = self.doc.add_paragraph()
        h2.add_run("2. CRONOLOGIA DAS REUNIÕES").bold = True
        for ata_dict in lista_atas:
            ata = ata_dict.get("ata", {})
            meta = ata.get("metadata", {})
            data = meta.get("data", "?")
            n_temas = len(ata.get("temas", []))
            n_enc = len(ata.get("encaminhamentos", []))
            self.doc.add_paragraph(
                f"– {data}: {n_temas} deliberação(ões), {n_enc} encaminhamento(s)"
            )
        self.doc.add_paragraph()

        # 3. Consolidação Temática
        h3 = self.doc.add_paragraph()
        h3.add_run("3. CONSOLIDAÇÃO TEMÁTICA").bold = True
        table = self.doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for cell, label in zip(hdr, ["Tema", "Deliberação", "Responsável", "Prazo"]):
            cell.text = label
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for ata_dict in lista_atas:
            for tema in ata_dict.get("ata", {}).get("temas", []):
                row = table.add_row().cells
                row[0].text = tema.get("titulo", "")[:80]
                row[1].text = tema.get("deliberacao", "")[:80]
                row[2].text = tema.get("relator", "")
                row[3].text = ""
        self.doc.add_paragraph()

        # 4. Encaminhamentos Pendentes
        h4 = self.doc.add_paragraph()
        h4.add_run("4. ENCAMINHAMENTOS PENDENTES").bold = True
        for ata_dict in lista_atas:
            for enc in ata_dict.get("ata", {}).get("encaminhamentos", []):
                if enc.get("status", "Pendente") == "Pendente":
                    self.doc.add_paragraph(
                        f"– {enc.get('descricao', '')} | "
                        f"Resp.: {enc.get('responsavel', '')} | "
                        f"Prazo: {enc.get('prazo', '')}"
                    )
        self.doc.add_paragraph()

        # 5. Recomendações
        h5 = self.doc.add_paragraph()
        h5.add_run("5. RECOMENDAÇÕES PARA INSTÂNCIAS SUPERIORES").bold = True
        self.doc.add_paragraph("1. [Recomendação 1 — a ser preenchida manualmente]")
        self.doc.add_paragraph("2. [Recomendação 2 — a ser preenchida manualmente]")
        self.doc.add_paragraph("3. [Recomendação 3 — a ser preenchida manualmente]")
        self.doc.add_paragraph()

        # 6. Conclusão
        h6 = self.doc.add_paragraph()
        h6.add_run("6. CONCLUSÃO").bold = True
        self.doc.add_paragraph("[Síntese estratégica e próximos passos — a ser preenchida manualmente]")

    # ------------------------------------------------------------------
    # SALVAR
    # ------------------------------------------------------------------

    def salvar(self, caminho: str) -> str:
        self.doc.save(caminho)
        return caminho


if __name__ == "__main__":
    print("DocxGenerator importado com sucesso.")
    print(f"python-docx disponível: {DOCX_AVAILABLE}")
